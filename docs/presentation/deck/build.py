# /// script
# requires-python = ">=3.11"
# dependencies = ["playwright>=1.50", "python-pptx>=1.0", "pillow>=11", "segno>=1.6", "pymupdf>=1.24"]
# ///
"""Builds the deck around slides.html (which itself is the animated deck you present from the browser):

    uv run docs/presentation/deck/build.py [--sheet PATH] [--no-video]

1. img/: the live-stand screenshots of shots/ as JPEG, and the QR of the stand; data/geo.js for the animations, with
   the walk of the live answer hero_walk for slide 3.
2. Checks that every number of slides_data.numbers() is written on the slides and that every photo of photos/ the
   slides show is credited (author and licence) on the sources slide.
3. out/slide_NN.png: the final frame of every slide; out/slide_NN.mp4: every slide with a map or counter animation
   (data-anim), 30 fps.
4. docs/presentation/deck.pdf (static final frames), deck.pptx (the frames, the animated slides as autoplaying
   looping videos, the speaker notes), speaker_notes.md, and a contact sheet with --sheet. Every PDF page is compared
   with the final frame of its slide: print lays out pages on its own, and a picture that crosses the page edge is
   moved to the next page there, so a missing phone or photo shows up as a page that differs from its frame.
The photos themselves are made once by photos/fetch.py. Needs ffmpeg on PATH and Playwright's Chromium
(`uv run --with playwright playwright install chromium` once).
"""

from __future__ import annotations

import argparse
import html as htmllib
import io
import json
import math
import re
import shutil
import subprocess
import sys
from pathlib import Path

from PIL import Image, ImageChops, ImageStat
from playwright.sync_api import sync_playwright

DECK = Path(__file__).resolve().parent
sys.path.insert(0, str(DECK))

import figures  # noqa: E402
import slides_data  # noqa: E402

PRESENTATION = DECK.parent
OUT = DECK / "out"
SHOTS = DECK / "shots"
IMG = DECK / "img"
PHOTOS = DECK / "photos"
W, H, FPS, HOLD_S = 1920, 1080, 30, 1.5


def prepare(src: str) -> None:
    IMG.mkdir(exist_ok=True)
    used = set(re.findall(r'src="img/([a-z0-9_]+)\.jpg"', src))
    for name in used:
        Image.open(SHOTS / f"{name}.png").convert("RGB").save(IMG / f"{name}.jpg", "JPEG", quality=88, optimize=True,
                                                               progressive=True)
    for stale in IMG.glob("*.jpg"):
        if stale.stem not in used:
            stale.unlink()
    import segno

    segno.make("https://anyq.an8kk.dev", error="m").save(str(IMG / "qr.svg"), scale=10, border=0, dark="#161B26")
    geo = figures.geo()
    geo["take"] = hero_take()
    (DECK / "data" / "geo.js").write_text(
        "window.GEO = " + json.dumps(geo, ensure_ascii=False, separators=(",", ":")) + ";\n", encoding="utf-8"
    )


# The stand's screenshot of that answer (shots/hero_walk_steps.png, 1920 px wide), measured once by the pixels of its
# markers: the centre of «Вы здесь» (the answer's origin) and of pin 1 («Японский сад»). Slide 3 lands its own map so
# that the stand's route covers ours: the pin of NU Park predicted from these two falls within 2 px of the screenshot's.
HERO_SHOT_W = 1920
HERO_HERE_PX = (742.5, 574.0)
HERO_PIN1_PX = (609.25, 384.4)


def hero_take() -> dict:
    """The walk slide 3 draws: the origin, the street line to the first place and the places of the live answer
    hero_walk (shots/shots.jsonl, place_legs), in the frame of the city map; and how that frame maps onto the
    stand's screenshot of the answer."""
    rows = [json.loads(x) for x in (SHOTS / "shots.jsonl").read_text(encoding="utf-8").splitlines()]
    answer = [r for r in rows if r["id"] == "hero_walk" and r.get("accepted")][-1]["answer"]

    def xy(lat: float, lon: float) -> list[float]:
        return [round(v, 2) for v in figures.CITY.xy(lon, lat)]

    places = answer["turn"]["places"]
    here, pin1 = xy(*answer["turn"]["origin"]), xy(places[0]["lat"], places[0]["lon"])
    scale = math.dist(HERO_HERE_PX, HERO_PIN1_PX) / math.dist(here, pin1)
    return {
        "start": here,
        "line": [xy(la, lo) for la, lo in answer["place_legs"][0]["shape"]],
        "places": [{"xy": xy(p["lat"], p["lon"]), "label": f"{p['name']} · {p['walk']['minutes']} мин"} for p in places],
        "match": {"width": HERO_SHOT_W, "here": HERO_HERE_PX, "scale": round(scale, 4)},
    }


def text_of(s: str) -> str:
    return re.sub(r"\s+", " ", htmllib.unescape(re.sub(r"<[^>]+>", " ", s))).strip()


def check_numbers(src: str) -> None:
    flat = re.sub(r"[\s\u00a0\u202f\u2009]", "", text_of(src))
    missing = [f"{k}: {v}" for k, v in slides_data.numbers().items() if re.sub(r"\s", "", v) not in flat]
    if missing:
        raise SystemExit("numbers not on the slides:\n  " + "\n  ".join(missing))


def check_credits(src: str) -> None:
    """CC BY-SA asks for the author, the licence and the link: each photo the slides show has them on the sources slide."""
    credits = json.loads((PHOTOS / "credits.json").read_text(encoding="utf-8"))
    sources = text_of(src[src.index('<section class="slide sources"'):])
    bad = []
    for name in sorted(set(re.findall(r'src="photos/([a-z0-9_]+)\.jpg"', src))):
        c = credits.get(name)
        if not c:
            bad.append(f"{name}: no entry in photos/credits.json")
            continue
        page = c["page"].removeprefix("https://")
        for need in (c["author"], c["license"], page):
            if need not in sources:
                bad.append(f"{name}: «{need}» is not on the sources slide")
    if bad:
        raise SystemExit("photo credits:\n  " + "\n  ".join(bad))


OVERFLOW_JS = """() => {
  const out = [];
  document.querySelectorAll('section.slide').forEach((s, i) => {
    const r = s.getBoundingClientRect();
    s.querySelectorAll('h1, h2, p, li, figure, .wordmark, .url, .num, .facts > div, .rk-row, .chain, .keys, .card, .call, .d-lab')
      .forEach((el) => {
        const b = el.getBoundingClientRect();
        if (!b.width || !b.height) return;
        if (b.right > r.right + 1 || b.bottom > r.bottom + 1 || b.left < r.left - 1 || b.top < r.top - 1)
          out.push(`slide ${i + 1}: ${el.tagName.toLowerCase()}.${el.className} leaves the slide`);
      });
  });
  return out;
}"""


def render(sheet: Path | None, video: bool) -> None:
    src = (DECK / "slides.html").read_text(encoding="utf-8")
    prepare(src)
    check_numbers(src)
    check_credits(src)
    sections = re.findall(r'<section class="slide[^"]*"[^>]*data-title="([^"]+)"[^>]*>(.*?)</section>', src, flags=re.S)
    titles = [text_of(t) for t, _ in sections]
    notes = [text_of(m.group(1)) if (m := re.search(r'<aside class="notes">(.*?)</aside>', body, re.S)) else ""
             for _, body in sections]
    srcs = [text_of(m.group(1)) if (m := re.search(r'<aside class="src">(.*?)</aside>', body, re.S)) else ""
            for _, body in sections]
    OUT.mkdir(exist_ok=True)
    for old in OUT.glob("slide_*"):
        old.unlink()
    frames = OUT / "frames"
    url = (DECK / "slides.html").as_uri()
    pngs: list[Path] = []
    movies: dict[int, Path] = {}
    with sync_playwright() as pw:
        browser = pw.chromium.launch()
        page = browser.new_page(viewport={"width": W, "height": H})
        page.goto(url + "?frame")
        page.wait_for_function("window.READY === true", timeout=60000)
        n = page.evaluate("document.querySelectorAll('section.slide').length")
        for i in range(n):
            d = page.evaluate(f"window.slideDur({i})")
            page.evaluate(f"window.seek({i}, 1e4)")
            png = OUT / f"slide_{i + 1:02d}.png"
            page.screenshot(path=str(png))
            pngs.append(png)
            if not video or not page.evaluate(f"window.animated({i})"):
                continue
            shutil.rmtree(frames, ignore_errors=True)
            frames.mkdir()
            for f in range(int(d * FPS) + 1):
                page.evaluate(f"window.seek({i}, {f / FPS})")
                page.screenshot(path=str(frames / f"f{f:04d}.jpg"), type="jpeg", quality=92)
            mp4 = OUT / f"slide_{i + 1:02d}.mp4"
            subprocess.run(
                ["ffmpeg", "-loglevel", "error", "-y", "-framerate", str(FPS), "-i", str(frames / "f%04d.jpg"),
                 "-vf", f"tpad=stop_mode=clone:stop_duration={HOLD_S},format=yuv420p", "-c:v", "libx264",
                 "-preset", "medium", "-crf", "21", "-movflags", "+faststart", str(mp4)],
                check=True,
            )
            movies[i] = mp4
        shutil.rmtree(frames, ignore_errors=True)
        page.goto(url + "?print")
        page.wait_for_function("window.READY === true", timeout=60000)
        for line in page.evaluate(OVERFLOW_JS):
            print("overflow:", line)
        page.pdf(path=str(PRESENTATION / "deck.pdf"), width=f"{W}px", height=f"{H}px", print_background=True,
                 margin={"top": "0", "right": "0", "bottom": "0", "left": "0"})
        browser.close()
    if not len(notes) == len(titles) == len(pngs) or not all(notes):
        raise SystemExit(f"{len(pngs)} slides, {len(titles)} titles, {sum(map(bool, notes))} speaker notes")
    check_pdf(pngs)
    build_pptx(pngs, movies, [f"{n}\n\nИсточники: {s}" if s else n for n, s in zip(notes, srcs)])
    write_notes(notes, srcs, titles)
    if sheet:
        contact_sheet(pngs, sheet)
    print(f"{len(pngs)} slides, {len(movies)} animated; deck.pdf {(PRESENTATION / 'deck.pdf').stat().st_size / 1e6:.2f} MB; "
          f"deck.pptx {(PRESENTATION / 'deck.pptx').stat().st_size / 1e6:.2f} MB; "
          f"{sum(len(n.split()) for n in notes)} words of speech")


def check_pdf(pngs: list[Path], limit: float = 6.0) -> None:
    """Each page of deck.pdf against the final frame of its slide, at 480 x 270: mean absolute difference of the pixels."""
    import pymupdf

    doc = pymupdf.open(PRESENTATION / "deck.pdf")
    if doc.page_count != len(pngs):
        raise SystemExit(f"deck.pdf has {doc.page_count} pages for {len(pngs)} slides")
    bad = []
    for i, (page, png) in enumerate(zip(doc, pngs), 1):
        pix = page.get_pixmap(dpi=96)
        a = Image.frombytes("RGB", (pix.width, pix.height), pix.samples).resize((480, 270), Image.LANCZOS)
        b = Image.open(png).convert("RGB").resize((480, 270), Image.LANCZOS)
        diff = sum(ImageStat.Stat(ImageChops.difference(a, b)).mean) / 3
        if diff > limit:
            bad.append(f"page {i}: differs from its frame by {diff:.1f}")
    if bad:
        raise SystemExit("deck.pdf:\n  " + "\n  ".join(bad))


# Start the movie with the slide and loop it: the timing PowerPoint writes for «Start: Automatically» and
# «Loop until Stopped» (python-pptx's own add_movie timing plays on click only).
TIMING = (
    '<p:timing xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"><p:tnLst><p:par>'
    '<p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
    '<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>'
    '<p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/>'
    '<p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond></p:stCondLst><p:childTnLst>'
    '<p:par><p:cTn id="4" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
    '<p:par><p:cTn id="5" presetID="1" presetClass="mediacall" presetSubtype="0" fill="hold" nodeType="afterEffect">'
    '<p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
    '<p:cmd type="call" cmd="playFrom(0.0)"><p:cBhvr><p:cTn id="6" dur="{ms}" fill="hold"/>'
    '<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:cmd>'
    '</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>'
    '</p:childTnLst></p:cTn>'
    '<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
    '<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>'
    '<p:video><p:cMediaNode vol="80000"><p:cTn id="7" repeatCount="indefinite" fill="hold" display="0">'
    '<p:stCondLst><p:cond delay="indefinite"/></p:stCondLst></p:cTn>'
    '<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cMediaNode></p:video>'
    '</p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>'
)


def build_pptx(pngs: list[Path], movies: dict[int, Path], notes: list[str]) -> None:
    from lxml import etree
    from pptx import Presentation
    from pptx.util import Emu

    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    blank = prs.slide_layouts[6]
    for i, (png, note) in enumerate(zip(pngs, notes)):
        s = prs.slides.add_slide(blank)
        buf = io.BytesIO()
        Image.open(png).convert("RGB").save(buf, "JPEG", quality=90, optimize=True)
        buf.seek(0)
        s.shapes.add_picture(buf, 0, 0, prs.slide_width, prs.slide_height)
        if i in movies:
            poster = io.BytesIO(buf.getvalue())
            mv = s.shapes.add_movie(str(movies[i]), 0, 0, prs.slide_width, prs.slide_height,
                                    poster_frame_image=poster, mime_type="video/mp4")
            secs = float(subprocess.run(["ffprobe", "-v", "error", "-show_entries", "format=duration", "-of",
                                         "csv=p=0", str(movies[i])], capture_output=True, text=True).stdout or 5)
            sld = s._element
            for old in sld.findall("{http://schemas.openxmlformats.org/presentationml/2006/main}timing"):
                sld.remove(old)
            sld.append(etree.fromstring(TIMING.format(spid=mv.shape_id, ms=int(secs * 1000))))
        s.notes_slide.notes_text_frame.text = note
    prs.save(PRESENTATION / "deck.pptx")


def write_notes(notes: list[str], srcs: list[str], titles: list[str]) -> None:
    lines = [
        "# Анық: текст к слайдам",
        "",
        "Что говорить на каждом слайде, 5–7 минут на всё; под текстом — откуда его факты. Источник файла:"
        " `docs/presentation/deck/slides.html` (блоки `notes` и `src`), пересборка `uv run docs/presentation/deck/build.py`.",
        "",
    ]
    for i, (t, n, s) in enumerate(zip(titles, notes, srcs), 1):
        lines += [f"## {i:02d}. {t}", "", n, ""]
        if s:
            lines += [f"Источники: {s}", ""]
    (PRESENTATION / "speaker_notes.md").write_text("\n".join(lines), encoding="utf-8")


def contact_sheet(pngs: list[Path], path: Path, cols: int = 3, tw: int = 640) -> None:
    th = tw * H // W
    gap = 16
    rows = (len(pngs) + cols - 1) // cols
    sheet = Image.new("RGB", (cols * tw + (cols + 1) * gap, rows * th + (rows + 1) * gap), "#3A3833")
    for i, p in enumerate(pngs):
        im = Image.open(p).convert("RGB").resize((tw, th), Image.LANCZOS)
        r, c = divmod(i, cols)
        sheet.paste(im, (gap + c * (tw + gap), gap + r * (th + gap)))
    path.parent.mkdir(parents=True, exist_ok=True)
    sheet.save(path, optimize=True)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--sheet", type=Path, help="write a contact sheet of all slides here")
    ap.add_argument("--no-video", action="store_true", help="skip the per-slide videos (and so the movies in the PPTX)")
    a = ap.parse_args()
    render(a.sheet, not a.no_video)


if __name__ == "__main__":
    main()
